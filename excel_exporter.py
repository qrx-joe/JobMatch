"""
Excel 导出器 - 将筛选结果导出为 Excel
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class ExcelExporter:
    def __init__(self, config_path=None):
        self.config_path = config_path

    def export(self, jobs, output_path):
        """导出岗位数据到Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "筛选结果"

        # 表头
        headers = ["所属地市", "服务单位", "岗位类型", "专业要求", "学历要求", "学位要求", "招募人数", "竞争比", "匹配度", "匹配说明"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")

        # 数据行
        for row_idx, job in enumerate(jobs, 2):
            ws.cell(row=row_idx, column=1, value=job.sheet_name)
            ws.cell(row=row_idx, column=2, value=job.unit)
            ws.cell(row=row_idx, column=3, value=job.job_type)
            ws.cell(row=row_idx, column=4, value=job.major)
            ws.cell(row=row_idx, column=5, value=job.education)
            ws.cell(row=row_idx, column=6, value=job.degree)
            ws.cell(row=row_idx, column=7, value=job.recruit_count)
            ws.cell(row=row_idx, column=8, value=job.competition_ratio if job.competition_ratio else 0)
            ws.cell(row=row_idx, column=9, value=job.match_level.value if hasattr(job.match_level, 'value') else str(job.match_level))
            ws.cell(row=row_idx, column=10, value=", ".join(job.match_reasons) if job.match_reasons else "")

        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        wb.save(output_path)
        return output_path