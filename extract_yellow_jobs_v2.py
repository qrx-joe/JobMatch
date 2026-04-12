#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
提取标黄的岗位，并与统计数据对照，生成Excel报告（改进版）
"""
import pandas as pd
import openpyxl
import glob
import os
import warnings
import sys
import io
warnings.filterwarnings('ignore')

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# 文件路径
dir_path = r'D:\EdgeDownload\QQ音乐\three-zhi-one-fu'
xlsx_files = [f for f in glob.glob(dir_path + '\\*.xlsx')
              if not os.path.basename(f).startswith('~$')
              and '标黄岗位统计信息' not in f]

if not xlsx_files:
    print("未找到xlsx文件")
    sys.exit(1)

jobs_file = xlsx_files[0]
stats_file = r'D:\EdgeDownload\QQ音乐\three-zhi-one-fu\P020260410598205300172.xls'

print(f"处理岗位文件: {os.path.basename(jobs_file)}")
print(f"处理统计文件: {os.path.basename(stats_file)}")

# ========== 第一步：读取统计文件 ==========
print("\n" + "="*80)
print("第一步：读取统计数据")
print("="*80)

stats_df = pd.read_excel(stats_file, sheet_name='sheet1', header=1)
stats_df.columns = ['服务单位', '岗位类型', '招募人数', '填报信息人数', '初审通过人数', '缴费人数']

# 清理服务单位名称，用于匹配
stats_map = {}
for idx, row in stats_df.iterrows():
    if pd.notna(row['服务单位']):
        key = str(row['服务单位']).strip()
        stats_map[key] = {
            '填报信息人数': int(row['填报信息人数']) if pd.notna(row['填报信息人数']) else 0,
            '初审通过人数': int(row['初审通过人数']) if pd.notna(row['初审通过人数']) else 0,
            '缴费人数': int(row['缴费人数']) if pd.notna(row['缴费人数']) else 0,
            '招募人数': int(row['招募人数']) if pd.notna(row['招募人数']) else 0
        }

print(f"统计文件共 {len(stats_map)} 条记录")

# ========== 第二步：读取岗位文件并识别黄色背景 ==========
print("\n" + "="*80)
print("第二步：识别标黄的岗位")
print("="*80)

# 使用openpyxl读取以获取样式
wb = openpyxl.load_workbook(jobs_file)

yellow_jobs = []

for sheet_name in wb.sheetnames:
    if sheet_name == '省林草局':  # 跳过省林草局
        continue

    print(f"\n检查Sheet: {sheet_name}")

    # 先用pandas读取该sheet获取正确的表格结构
    try:
        pd_df = pd.read_excel(jobs_file, sheet_name=sheet_name, header=None)
    except Exception as e:
        print(f"  读取失败: {e}")
        continue

    if len(pd_df) < 4:
        print(f"  数据行数不足")
        continue

    # 获取表头行
    header_row1 = pd_df.iloc[1].tolist()  # 主表头
    header_row2 = pd_df.iloc[2].tolist()  # 子表头

    # 合并表头
    headers = []
    for i, (h1, h2) in enumerate(zip(header_row1, header_row2)):
        h1_str = str(h1) if pd.notna(h1) else ''
        h2_str = str(h2) if pd.notna(h2) else ''
        if h2_str and h2_str != 'nan':
            headers.append(f"{h1_str}_{h2_str}")
        else:
            headers.append(h1_str)

    # 使用openpyxl获取黄色背景的行号
    ws = wb[sheet_name]
    yellow_row_indices = set()

    for row_idx, row in enumerate(ws.iter_rows(min_row=4, max_row=ws.max_row), start=4):
        first_cell = row[0] if len(row) > 0 else None
        if first_cell and first_cell.value is not None:
            fill = first_cell.fill
            is_yellow = False
            if fill and fill.patternType == 'solid' and fill.fgColor:
                color = fill.fgColor
                if color.rgb and 'FFFF00' in str(color.rgb).upper():
                    is_yellow = True
                elif color.rgb and 'FF00' in str(color.rgb).upper() and str(color.rgb).upper().startswith('FF'):
                    is_yellow = True
                elif color.index == 6 or color.index == '6':
                    is_yellow = True
            if is_yellow:
                # pandas行索引从0开始，excel从1开始，数据从第4行开始（索引3）
                pandas_row_idx = row_idx - 1
                yellow_row_indices.add(pandas_row_idx)

    print(f"  找到 {len(yellow_row_indices)} 个标黄岗位")

    # 提取标黄行的数据
    for row_idx in yellow_row_indices:
        if row_idx < len(pd_df):
            row_data = pd_df.iloc[row_idx].tolist()

            # 构建岗位信息字典
            job = {
                '所属地市': sheet_name,
                '行号': row_idx + 1,
                '序号': row_data[0] if len(row_data) > 0 and pd.notna(row_data[0]) else '',
                '服务单位': row_data[1] if len(row_data) > 1 and pd.notna(row_data[1]) else '',
                '岗位类型': row_data[2] if len(row_data) > 2 and pd.notna(row_data[2]) else '',
                '服务类别': row_data[3] if len(row_data) > 3 and pd.notna(row_data[3]) else '',
                '招募人数': row_data[4] if len(row_data) > 4 and pd.notna(row_data[4]) else '',
                '学历': row_data[5] if len(row_data) > 5 and pd.notna(row_data[5]) else '',
                '学位': row_data[6] if len(row_data) > 6 and pd.notna(row_data[6]) else '',
                '专业': row_data[7] if len(row_data) > 7 and pd.notna(row_data[7]) else '',
                '相关资格': row_data[8] if len(row_data) > 8 and pd.notna(row_data[8]) else '',
                '其他': row_data[9] if len(row_data) > 9 and pd.notna(row_data[9]) else '',
                '联系电话': row_data[10] if len(row_data) > 10 and pd.notna(row_data[10]) else '',
                '联系人': row_data[11] if len(row_data) > 11 and pd.notna(row_data[11]) else '',
                '岗位描述': row_data[12] if len(row_data) > 12 and pd.notna(row_data[12]) else '',
                '福利待遇': row_data[13] if len(row_data) > 13 and pd.notna(row_data[13]) else ''
            }
            yellow_jobs.append(job)

print(f"\n总共找到 {len(yellow_jobs)} 个标黄岗位")

# ========== 第三步：对照统计文件，补充数据 ==========
print("\n" + "="*80)
print("第三步：对照统计数据")
print("="*80)

# 尝试匹配统计数据
for job in yellow_jobs:
    service_unit = str(job['服务单位']).strip()
    job_type = str(job['岗位类型']).strip()

    matched = False
    best_match_key = None
    best_match_score = 0

    for stats_key in stats_map.keys():
        # 计算匹配分数
        score = 0
        if service_unit in stats_key:
            score += 10
        if stats_key in service_unit:
            score += 5

        # 岗位类型匹配
        if job_type in stats_key:
            score += 3

        # 提取单位名称部分匹配
        if '-' in stats_key:
            unit_parts = stats_key.split('-')
            for part in unit_parts:
                if part in service_unit or service_unit in part:
                    score += 2

        if score > best_match_score:
            best_match_score = score
            best_match_key = stats_key

    if best_match_key and best_match_score >= 5:
        job['填报信息人数'] = stats_map[best_match_key]['填报信息人数']
        job['初审通过人数'] = stats_map[best_match_key]['初审通过人数']
        job['缴费人数'] = stats_map[best_match_key]['缴费人数']
        job['统计招募人数'] = stats_map[best_match_key]['招募人数']
        job['匹配状态'] = '已匹配' if best_match_score >= 10 else '部分匹配'
        job['匹配键'] = best_match_key
        matched = True

    if not matched:
        job['填报信息人数'] = 0
        job['初审通过人数'] = 0
        job['缴费人数'] = 0
        job['统计招募人数'] = 0
        job['匹配状态'] = '未匹配'
        job['匹配键'] = ''

matched_count = sum(1 for job in yellow_jobs if job['匹配状态'] == '已匹配')
partial_count = sum(1 for job in yellow_jobs if job['匹配状态'] == '部分匹配')
unmatched_count = sum(1 for job in yellow_jobs if job['匹配状态'] == '未匹配')

print(f"匹配结果：已匹配 {matched_count} 条，部分匹配 {partial_count} 条，未匹配 {unmatched_count} 条")

# 显示未匹配的岗位
if unmatched_count > 0:
    print("\n未匹配的岗位:")
    for job in yellow_jobs:
        if job['匹配状态'] == '未匹配':
            print(f"  - {job['所属地市']}: {job['服务单位']} ({job['岗位类型']})")

# ========== 第四步：生成Excel文件 ==========
print("\n" + "="*80)
print("第四步：生成Excel文件")
print("="*80)

# 准备输出数据
output_data = []
for job in yellow_jobs:
    output_data.append({
        '所属地市': job['所属地市'],
        '序号': job['序号'],
        '服务单位': job['服务单位'],
        '岗位类型': job['岗位类型'],
        '服务类别': job['服务类别'],
        '招募人数': job['招募人数'],
        '学历': job['学历'],
        '学位': job['学位'],
        '专业': job['专业'],
        '相关资格': job['相关资格'],
        '其他': job['其他'],
        '联系电话': job['联系电话'],
        '联系人': job['联系人'],
        '岗位描述': job['岗位描述'],
        '福利待遇': job['福利待遇'],
        '填报信息人数': job['填报信息人数'],
        '初审通过人数': job['初审通过人数'],
        '缴费人数': job['缴费人数'],
        '匹配状态': job['匹配状态']
    })

# 创建DataFrame并保存
output_df = pd.DataFrame(output_data)

output_file = r'D:\EdgeDownload\QQ音乐\three-zhi-one-fu\标黄岗位统计信息.xlsx'

# 使用openpyxl引擎保存，并设置列宽
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    output_df.to_excel(writer, index=False, sheet_name='标黄岗位')

    # 获取工作表并调整列宽
    worksheet = writer.sheets['标黄岗位']
    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            except:
                pass
        adjusted_width = min(max_length + 2, 50)  # 最大50
        worksheet.column_dimensions[column_letter].width = adjusted_width

print(f"\n✓ Excel文件已生成: {output_file}")
print(f"  共 {len(output_data)} 条记录")

# 统计摘要
total_recruit = sum(int(job['招募人数']) if str(job['招募人数']).isdigit() else 0 for job in yellow_jobs)
total_apply = sum(job['填报信息人数'] for job in yellow_jobs)
total_review = sum(job['初审通过人数'] for job in yellow_jobs)
total_pay = sum(job['缴费人数'] for job in yellow_jobs)

print(f"\n统计摘要:")
print(f"  总招募人数: {total_recruit}")
print(f"  总填报信息人数: {total_apply}")
print(f"  总初审通过人数: {total_review}")
print(f"  总缴费人数: {total_pay}")
if total_recruit > 0:
    print(f"  平均竞争比: {total_pay/total_recruit:.2f} : 1")
