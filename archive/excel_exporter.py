#!/usr/bin/env python3
"""
Excel导出器 - 生成带颜色标记的筛选结果
"""

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from job_matcher import Job, MatchLevel


class ExcelExporter:
    """Excel导出器"""

    # 颜色定义
    COLORS = {
        "perfect": "C6EFCE",  # 浅绿色
        "perfect_text": "006100",
        "partial": "FFEB9C",  # 浅黄色
        "partial_text": "9C5700",
        "mismatch": "FFC7CE",  # 浅红色
        "mismatch_text": "9C0006",
        "header": "4472C4",  # 蓝色表头
        "header_text": "FFFFFF",
    }

    def __init__(self, config_path: str = "config.yaml"):
        with open(config_path, encoding="utf-8") as f:
            self.config = yaml.safe_load(f)

    def export(self, jobs: list[Job], output_path: str):
        """导出为Excel文件"""
        # 准备数据
        output_mode = self.config.get("output", {}).get("输出模式", "全量标记")
        sort_by = self.config.get("output", {}).get("排序方式", "匹配度降序")

        # 筛选数据
        if output_mode == "match_only":
            export_jobs = [j for j in jobs if j.match_level != MatchLevel.MISMATCH]
        else:
            export_jobs = jobs

        # 排序
        if sort_by == "匹配度降序":
            export_jobs.sort(key=lambda j: j.match_score, reverse=True)
        elif sort_by == "竞争比升序":
            export_jobs.sort(key=lambda j: j.competition_ratio if j.paid > 0 else float("inf"))
        elif sort_by == "竞争比降序":
            export_jobs.sort(key=lambda j: j.competition_ratio if j.paid > 0 else 0, reverse=True)

        # 创建DataFrame
        data = []
        for job in export_jobs:
            row = {
                "匹配度": job.match_level.value,
                "匹配分数": job.match_score,
                "所属地市": job.sheet_name,
                "序号": job.index,
                "服务单位": job.unit,
                "岗位类型": job.job_type,
                "服务类别": job.service_category,
                "招募人数": job.recruit_count,
                "学历要求": job.education,
                "学位要求": job.degree,
                "专业要求": job.major,
                "相关资格": job.qualifications,
                "其他要求": job.other,
                "联系电话": job.phone,
                "联系人": job.contact,
                "填报人数": job.applicants,
                "缴费人数": job.paid,
                "竞争比": f"{job.competition_ratio:.1f}:1" if job.paid > 0 else "暂无数据",
                "竞争比数值": job.competition_ratio,
                "匹配说明": "；".join(job.match_reasons),
                "不匹配原因": "；".join(job.mismatch_reasons) if job.mismatch_reasons else "",
            }
            data.append(row)

        df = pd.DataFrame(data)

        # 创建工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "筛选结果"

        # 写入数据
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                # 表头样式
                if r_idx == 1:
                    cell.fill = PatternFill(
                        start_color=self.COLORS["header"],
                        end_color=self.COLORS["header"],
                        fill_type="solid",
                    )
                    cell.font = Font(bold=True, color=self.COLORS["header_text"])
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    # 数据行根据匹配度着色
                    match_level = data[r_idx - 2]["匹配度"]
                    if match_level == MatchLevel.PERFECT.value:
                        cell.fill = PatternFill(
                            start_color=self.COLORS["perfect"],
                            end_color=self.COLORS["perfect"],
                            fill_type="solid",
                        )
                    elif match_level == MatchLevel.PARTIAL.value:
                        cell.fill = PatternFill(
                            start_color=self.COLORS["partial"],
                            end_color=self.COLORS["partial"],
                            fill_type="solid",
                        )
                    else:
                        cell.fill = PatternFill(
                            start_color=self.COLORS["mismatch"],
                            end_color=self.COLORS["mismatch"],
                            fill_type="solid",
                        )

                    cell.alignment = Alignment(vertical="center", wrap_text=True)

        # 设置列宽
        column_widths = {
            "A": 10,  # 匹配度
            "B": 10,  # 匹配分数
            "C": 12,  # 所属地市
            "D": 8,  # 序号
            "E": 40,  # 服务单位
            "F": 12,  # 岗位类型
            "G": 15,  # 服务类别
            "H": 10,  # 招募人数
            "I": 15,  # 学历要求
            "J": 15,  # 学位要求
            "K": 30,  # 专业要求
            "L": 20,  # 相关资格
            "M": 20,  # 其他要求
            "N": 15,  # 联系电话
            "O": 10,  # 联系人
            "P": 10,  # 填报人数
            "Q": 10,  # 缴费人数
            "R": 12,  # 竞争比
            "S": 12,  # 竞争比数值（用于排序）
            "T": 40,  # 匹配说明
            "U": 40,  # 不匹配原因
        }

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # 设置行高
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 30

        # 冻结首行
        ws.freeze_panes = "A2"

        # 添加筛选器
        ws.auto_filter.ref = ws.dimensions

        # 隐藏"竞争比数值"列（仅用于排序）
        ws.column_dimensions["S"].hidden = True

        # 保存
        wb.save(output_path)
        print(f"  已导出 {len(data)} 条记录到 {output_path}")


if __name__ == "__main__":
    # 测试
    from job_matcher import Job, MatchLevel

    test_jobs = [
        Job(
            sheet_name="太原市",
            unit="测试单位1",
            job_type="管理1",
            major="经济学类",
            education="本科及以上",
            match_level=MatchLevel.PERFECT,
            match_score=100,
            paid=50,
            recruit_count=1,
            competition_ratio=50.0,
        ),
        Job(
            sheet_name="吕梁市",
            unit="测试单位2",
            job_type="专技1",
            major="不限",
            education="本科及以上",
            match_level=MatchLevel.PARTIAL,
            match_score=80,
            paid=30,
            recruit_count=1,
            competition_ratio=30.0,
        ),
        Job(
            sheet_name="大同市",
            unit="测试单位3",
            job_type="管理2",
            major="法学",
            education="研究生",
            match_level=MatchLevel.MISMATCH,
            match_score=0,
            paid=100,
            recruit_count=1,
            competition_ratio=100.0,
        ),
    ]

    exporter = ExcelExporter()
    exporter.export(test_jobs, "test_output.xlsx")
    print("测试导出完成: test_output.xlsx")
